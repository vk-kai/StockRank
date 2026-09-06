# -*- coding: utf-8 -*-
"""小程序数据后台管理接口测试:鉴权(仅vk登录)、趋势聚合、成绩/计数/房间增删改查。"""
import json
import os
import tempfile
import time
import unittest
from datetime import datetime, timedelta, timezone

from flask import Flask

from routes import mp_game_routes as m
from routes import mp_sec_routes as sec
from routes.mp_admin_routes import mp_admin_bp


def _utcnow_str():
    return datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')


class MpAdminTestCase(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.mkdtemp()
        m.DB_FILE = os.path.join(tmp, 'game.db')
        m._rl_store.clear()
        m._rank_cache.clear()
        m._overall_cache['ts'] = 0.0
        m._overall_cache['board'] = None
        m._stat_cache.clear()
        m._stat_dedup.clear()
        m._nick_ok_cache.clear()
        m._last_purge = 0.0
        # 游戏接口 auth_key 置空(集成测试用)
        sec.MP_SEC_CONFIG_FILE = os.path.join(tmp, 'sec.json')
        with open(sec.MP_SEC_CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump({'auth_key': ''}, f)

        app = Flask(__name__)
        app.secret_key = 'test'
        app.register_blueprint(mp_admin_bp)
        app.register_blueprint(m.mp_game_bp)
        self.client = app.test_client()

    def _login(self):
        with self.client.session_transaction() as s:
            s['stockrank_user'] = 'vk'

    def _add_score(self, quiz='test_a', openid='oA', nickname='甲', score=80,
                   full=100, dur=60000):
        with m._db() as conn:
            conn.execute('''INSERT INTO scores
                (quiz_id, openid, nickname, score, full_score, duration_ms, answers, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                (quiz, openid, nickname, score, full, dur, None, _utcnow_str()))


class AuthTests(MpAdminTestCase):
    def test_no_session_401(self):
        for url in ('/api/mp-admin/overview', '/api/mp-admin/scores',
                    '/api/mp-admin/quiz-rank?quiz_id=q',
                    '/api/mp-admin/stats', '/api/mp-admin/rooms'):
            r = self.client.get(url)
            self.assertEqual(r.status_code, 401, msg=url)
        r = self.client.post('/api/mp-admin/scores/delete', json={'password': 'pw'})
        self.assertEqual(r.status_code, 401)

    def test_write_ops_only_need_admin_session(self):
        # 二次动态密码已按需求移除(后台仅 vk 一人可见):登录即可写,
        # 不带 password 字段照常成功;body 里残留 password 也被忽略
        self._login()
        r = self.client.post('/api/mp-admin/scores/update',
                             json={'quiz_id': 'nope', 'openid': 'o', 'nickname': 'x'})
        self.assertEqual(r.status_code, 404)  # 已过鉴权,走到业务校验(记录不存在)
        r2 = self.client.post('/api/mp-admin/stats/save', json={'key': 'k', 'count': 1})
        self.assertEqual(r2.status_code, 200)
        self.assertTrue(r2.get_json()['success'])


class OverviewTests(MpAdminTestCase):
    def test_overview_shape_and_totals(self):
        self._login()
        self._add_score(quiz='test_a', openid='oA', score=80)
        self._add_score(quiz='test_b', openid='oB', score=90)
        with m._db() as conn:
            conn.execute("INSERT INTO stats(key, count) VALUES ('test_a', 5)")
            conn.execute("INSERT INTO stats(key, count) VALUES ('tool_foodwheel', 2)")
            conn.execute("INSERT INTO rooms(room_code, quiz_id, openid_a, state, expires_at, created_at)"
                         " VALUES ('1234', 'pk', 'oA', 'finished', '2099-01-01 00:00:00', ?)",
                         (_utcnow_str(),))
        r = self.client.get('/api/mp-admin/overview').get_json()
        self.assertTrue(r['success'])
        d = r['data']
        # 口径:tests 只含 test_* 计数,tool_usage 只含 tool_*,articles 只含 article_*
        self.assertEqual(d['totals']['tests'], 5)
        self.assertEqual(d['totals']['tool_usage'], 2)
        self.assertEqual(d['totals']['articles'], 0)
        self.assertEqual(d['totals']['article_count'], 0)
        self.assertEqual(d['totals']['players'], 2)
        self.assertEqual(d['totals']['quizzes'], 2)
        self.assertEqual(d['totals']['scores'], 2)
        self.assertEqual(d['totals']['rooms'], 1)
        self.assertEqual(d['totals']['rooms_finished'], 1)
        # 两榜拆分:测试榜/工具榜,带中文名(内置映射兜底,映射按小写key匹配)
        self.assertEqual(d['test_counts'],
                         [{'key': 'test_a', 'name': 'test_a', 'count': 5}])
        self.assertEqual(d['tool_counts'],
                         [{'key': 'tool_foodwheel', 'name': '吃什么转盘', 'count': 2}])
        self.assertEqual(d['article_counts'], [])
        # 近30天:今天应有 2 条新增成绩、1 个房间;无流水则计数趋势为0
        self.assertEqual(len(d['daily']['days']), 30)
        self.assertEqual(d['daily']['scores'][-1], 2)
        self.assertEqual(d['daily']['rooms'][-1], 1)
        self.assertEqual(sum(d['daily']['scores']), 2)
        self.assertEqual(sum(d['daily']['tests']), 0)
        self.assertEqual(sum(d['daily']['tools']), 0)
        self.assertEqual(sum(d['daily']['articles']), 0)
        # 热度排行:两测试并列,人数对;name 无映射时回退 quiz_id
        by_quiz = {q['quiz_id']: q['players'] for q in d['quiz_top']}
        self.assertEqual(by_quiz, {'test_a': 1, 'test_b': 1})
        for q in d['quiz_top']:
            self.assertEqual(q['name'], q['quiz_id'])
        # 无人参与≥3个测试 → 综合榜为空
        self.assertEqual(d['overall'], {'total': 0, 'top': []})
        # 今日按小时:24个桶,当天总和与按日口径的今天一致
        h = d['hourly']
        self.assertEqual(len(h['hours']), 24)
        self.assertEqual(h['hours'][0], '00')
        self.assertEqual(h['hours'][-1], '23')
        self.assertEqual(sum(h['tests']), 0)
        self.assertEqual(sum(h['tools']), 0)
        self.assertEqual(sum(h['articles']), 0)
        self.assertEqual(sum(h['scores']), d['daily']['scores'][-1])
        self.assertEqual(sum(h['rooms']), d['daily']['rooms'][-1])
        # 今日热力:当天无计数流水 → 空items/max=0;近30天热力同理
        self.assertEqual(d['hourly_top'], {'items': [], 'max': 0})
        self.assertEqual(d['daily_top'], {'items': [], 'max': 0})

    def test_name_priority_stored_over_builtin(self):
        # 名称优先级:小程序上报的 name > 内置映射 > 原样 key
        self._login()
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'test_mbti', 'openid': 'oA', 'name': '自定义名'})
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'tool_foodWheel', 'openid': 'oB'})  # 无name→内置(驼峰自动归一小写)
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'test_zzz', 'openid': 'oC'})          # 无映射→key
        r = self.client.get('/api/mp-admin/stats').get_json()
        names = {i['key']: i['name'] for i in r['data']['items']}
        self.assertEqual(names['test_mbti'], '自定义名')
        self.assertEqual(names['tool_foodwheel'], '吃什么转盘')
        self.assertEqual(names['test_zzz'], 'test_zzz')

    def test_overview_includes_overall_rank(self):
        # oA 参与3个测试达标上榜;oB 只测1个不计入综合榜
        self._login()
        self._add_score(quiz='quiz1', openid='oA', score=90)
        self._add_score(quiz='quiz1', openid='oB', score=60)
        self._add_score(quiz='quiz2', openid='oA', score=80)
        self._add_score(quiz='quiz3', openid='oA', score=70)
        d = self.client.get('/api/mp-admin/overview').get_json()['data']
        self.assertEqual(d['overall']['total'], 1)
        top = d['overall']['top']
        self.assertEqual(len(top), 1)
        # 击败率:quiz1 1/2、quiz2 0/1、quiz3 0/1 → 平均 50/3 = 16.7;昵称取该玩家最近一条成绩
        self.assertEqual(top[0], {'rank': 1, 'nickname': '甲',
                                  'avg_beat': 16.7, 'quizzes': 3})

    def test_stat_inc_feeds_daily_trend(self):
        # 集成:小程序计数接口真实+1后,趋势图当天人次应增加;工具key进工具曲线
        self._login()
        base_t = self.client.get('/api/mp-admin/overview').get_json()['data']['daily']['tests'][-1]
        base_g = self.client.get('/api/mp-admin/overview').get_json()['data']['daily']['tools'][-1]
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'test_careerFit', 'openid': 'oZ'})
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'tool_danmaku', 'openid': 'oZ'})
        d = self.client.get('/api/mp-admin/overview').get_json()['data']['daily']
        self.assertEqual(d['tests'][-1], base_t + 1)
        self.assertEqual(d['tools'][-1], base_g + 1)
        # 小时分布与按日今天值一致(inc 落在当前北京小时桶)
        h = self.client.get('/api/mp-admin/overview').get_json()['data']['hourly']
        self.assertEqual(sum(h['tests']), d['tests'][-1])
        self.assertEqual(sum(h['tools']), d['tools'][-1])
        self.assertTrue(any(c > 0 for c in h['tests']))
        # 今日热力:inc 过的两个 key 进榜,内置映射出中文名,各小时次数总和=当天次数
        top = self.client.get('/api/mp-admin/overview').get_json()['data']['hourly_top']
        items = {i['key']: i for i in top['items']}
        self.assertIn('test_careerfit', items)
        self.assertEqual(items['test_careerfit']['name'], '职业适配测试')
        self.assertEqual(len(items['test_careerfit']['hours']), 24)
        self.assertEqual(sum(items['test_careerfit']['hours']), 1)
        self.assertIn('tool_danmaku', items)
        self.assertEqual(items['tool_danmaku']['name'], '手持弹幕')
        self.assertGreaterEqual(top['max'], 1)
        # 近30天热力:同两个 key 进榜,days 30桶,最后一列(今天)=当天总次数
        dtop = self.client.get('/api/mp-admin/overview').get_json()['data']['daily_top']
        ditems = {i['key']: i for i in dtop['items']}
        self.assertIn('test_careerfit', ditems)
        self.assertEqual(ditems['test_careerfit']['name'], '职业适配测试')
        self.assertEqual(len(ditems['test_careerfit']['days']), 30)
        self.assertEqual(ditems['test_careerfit']['days'][-1], 1)
        self.assertEqual(sum(ditems['test_careerfit']['days']), 1)
        self.assertIn('tool_danmaku', ditems)
        self.assertEqual(ditems['tool_danmaku']['days'][-1], 1)
        self.assertGreaterEqual(dtop['max'], 1)
        # article_* 文章阅读:计数照常入库,但不进测试/工具热力图(量大口径不同)
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'article_intro', 'openid': 'oZ', 'name': '很' * 40})
        top2 = self.client.get('/api/mp-admin/overview').get_json()['data']['hourly_top']
        self.assertNotIn('article_intro', {i['key'] for i in top2['items']})
        dtop2 = self.client.get('/api/mp-admin/overview').get_json()['data']['daily_top']
        self.assertNotIn('article_intro', {i['key'] for i in dtop2['items']})
        # 长标题(>30字符,文章标题常见)照常入库;数据管理列表能看到
        items = self.client.get('/api/mp-admin/stats').get_json()['data']['items']
        art = next(i for i in items if i['key'] == 'article_intro')
        self.assertEqual(art['name'], '很' * 40)
        self.assertEqual(art['count'], 1)
        # 文章进 overview 的文章榜/趋势/总量(与测试/工具同接口,仅前缀不同)
        ov2 = self.client.get('/api/mp-admin/overview').get_json()['data']
        self.assertEqual(ov2['totals']['articles'], 1)
        self.assertEqual(ov2['totals']['article_count'], 1)
        acnts = {i['key']: i for i in ov2['article_counts']}
        self.assertEqual(acnts['article_intro']['name'], '很' * 40)
        self.assertEqual(acnts['article_intro']['count'], 1)
        self.assertEqual(ov2['daily']['articles'][-1], 1)
        self.assertEqual(sum(ov2['daily']['articles']), 1)
        self.assertEqual(sum(ov2['hourly']['articles']), 1)


class ExportTests(MpAdminTestCase):
    def test_export_requires_admin(self):
        r = self.client.get('/api/mp-admin/export')
        self.assertEqual(r.status_code, 401)

    def test_export_xlsx_sheets_and_numbers(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self._login()
        self._add_score(quiz='test_a', openid='oA', score=80)
        self._add_score(quiz='test_a', openid='oB', score=40)
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'test_a', 'openid': 'oA', 'name': '阿尔法测试'})
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'tool_foodwheel', 'openid': 'oA'})
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'article_intro', 'openid': 'oA', 'name': '开篇文章'})
        r = self.client.get('/api/mp-admin/export')
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r.headers['Content-Type'])
        self.assertIn('.xlsx', r.headers['Content-Disposition'])
        wb = load_workbook(BytesIO(r.data))
        self.assertEqual(wb.sheetnames,
                         ['概览', '测试明细', '工具明细', '文章明细', '每日趋势',
                          '今日按小时', '综合排名', 'PK房间'])
        # 测试明细:上报名生效,一条流水 + 两个去重参与者,得分率均值(80+40)/2
        rows = list(wb['测试明细'].iter_rows(values_only=True))
        head = list(rows[0])
        row = next(x for x in rows[1:] if x[0] == 'test_a')
        self.assertEqual(row[head.index('名称')], '阿尔法测试')
        self.assertEqual(row[head.index('全史完成次数')], 1)
        self.assertEqual(row[head.index('近30天次数')], 1)
        self.assertEqual(row[head.index('去重参与人数')], 2)
        self.assertEqual(row[head.index('平均得分率(%)')], 60.0)
        self.assertEqual(row[head.index('状态')], '活跃')
        # 工具明细:无上报名 → 内置中文名兜底
        trows = list(wb['工具明细'].iter_rows(values_only=True))
        thead = list(trows[0])
        trow = next(x for x in trows[1:] if x[0] == 'tool_foodwheel')
        self.assertEqual(trow[thead.index('名称')], '吃什么转盘')
        self.assertEqual(trow[thead.index('全史使用次数')], 1)
        # 文章明细:article_* 阅读计数单列一页,标题即上报名
        arows = list(wb['文章明细'].iter_rows(values_only=True))
        ahead = list(arows[0])
        arow = next(x for x in arows[1:] if x[0] == 'article_intro')
        self.assertEqual(arow[ahead.index('标题')], '开篇文章')
        self.assertEqual(arow[ahead.index('全史阅读次数')], 1)
        self.assertEqual(arow[ahead.index('状态')], '活跃')
        # 每日趋势:30行,最后一行=今天,当天测试次数=1
        drows = list(wb['每日趋势'].iter_rows(values_only=True))
        self.assertEqual(len(drows) - 1, 30)
        self.assertEqual(drows[-1][2], 1)
        # 概览:全史测试完成次数与去重人数
        ov = {x[0]: x[1] for x in wb['概览'].iter_rows(values_only=True) if x[0]}
        self.assertEqual(ov['测试完成总次数'], 1)
        self.assertEqual(ov['参与人数(玩过测试,去重)'], 2)
        self.assertEqual(ov['文章阅读总次数'], 1)


class ScoresTests(MpAdminTestCase):
    def test_list_pagination_and_filter(self):
        self._login()
        for i in range(25):
            self._add_score(quiz='test_a', openid=f'o{i:02d}')
        self._add_score(quiz='test_b', openid='oXX')
        r = self.client.get('/api/mp-admin/scores').get_json()
        self.assertEqual(r['data']['total'], 26)
        self.assertEqual(len(r['data']['items']), 20)
        self.assertEqual(r['data']['total_pages'], 2)
        r2 = self.client.get('/api/mp-admin/scores?quiz_id=test_b').get_json()
        self.assertEqual(r2['data']['total'], 1)
        self.assertEqual(r2['data']['items'][0]['quiz_id'], 'test_b')
        # 昵称为空时回退默认昵称
        with m._db() as conn:
            conn.execute('UPDATE scores SET nickname = NULL WHERE openid = ?', ('oXX',))
        r3 = self.client.get('/api/mp-admin/scores?quiz_id=test_b').get_json()
        self.assertEqual(r3['data']['items'][0]['nickname'], '匿名测试者')

    def test_update_and_delete(self):
        self._login()
        self._add_score(quiz='test_a', openid='oA', score=80)
        r = self.client.post('/api/mp-admin/scores/update', json={
            'password': 'pw', 'quiz_id': 'test_a', 'openid': 'oA',
            'nickname': '新名', 'score': 95, 'duration_ms': 30000,
        }).get_json()
        self.assertTrue(r['success'])
        with m._db() as conn:
            row = conn.execute("SELECT * FROM scores WHERE openid='oA'").fetchone()
        self.assertEqual(row['nickname'], '新名')
        self.assertEqual(row['score'], 95)
        self.assertEqual(row['duration_ms'], 30000)
        # 非法输入
        bad = self.client.post('/api/mp-admin/scores/update', json={
            'password': 'pw', 'quiz_id': 'test_a', 'openid': 'oA', 'score': -1})
        self.assertEqual(bad.status_code, 400)
        # 删除 + 404
        r2 = self.client.post('/api/mp-admin/scores/delete', json={
            'password': 'pw', 'quiz_id': 'test_a', 'openid': 'oA'}).get_json()
        self.assertTrue(r2['success'])
        r3 = self.client.post('/api/mp-admin/scores/delete', json={
            'password': 'pw', 'quiz_id': 'test_a', 'openid': 'oA'})
        self.assertEqual(r3.status_code, 404)

    def test_create(self):
        self._login()
        body = {'password': 'pw', 'quiz_id': 'test_a', 'openid': 'oNEW',
                'nickname': '补录', 'score': 66, 'full_score': 100, 'duration_ms': 45000}
        r = self.client.post('/api/mp-admin/scores/create', json=body).get_json()
        self.assertTrue(r['success'])
        with m._db() as conn:
            row = conn.execute("SELECT * FROM scores WHERE openid='oNEW'").fetchone()
        self.assertEqual(row['nickname'], '补录')
        self.assertEqual(row['score'], 66)
        self.assertTrue(row['created_at'])
        # 重复新增同一 (quiz_id, openid) → 409 提示走编辑
        r2 = self.client.post('/api/mp-admin/scores/create', json=body)
        self.assertEqual(r2.status_code, 409)
        # 参数校验:缺昵称 / 分数超满分 / 缺密码
        bad1 = self.client.post('/api/mp-admin/scores/create',
                                json={**body, 'openid': 'o2', 'nickname': ''})
        self.assertEqual(bad1.status_code, 400)
        bad2 = self.client.post('/api/mp-admin/scores/create',
                                json={**body, 'openid': 'o3', 'score': 101})
        self.assertEqual(bad2.status_code, 400)
        # 二次密码已移除:不带 password 字段照常入库
        no_pwd = dict(body, openid='o4')
        no_pwd.pop('password', None)
        ok4 = self.client.post('/api/mp-admin/scores/create', json=no_pwd)
        self.assertEqual(ok4.status_code, 200)
        # 未登录 401
        self.client.get('/api/mp-admin/logout')  # 无此路由也无妨,session 仍在
        with self.client.session_transaction() as s:
            s.clear()
        bad4 = self.client.post('/api/mp-admin/scores/create', json=body)
        self.assertEqual(bad4.status_code, 401)


class QuizRankTests(MpAdminTestCase):
    def test_rank_order_and_pagination(self):
        self._login()
        # 甲80分60秒、乙90分60秒、丙90分30秒:排序应为 丙>乙>甲(同分用时短优先)
        self._add_score(quiz='test_a', openid='oA', nickname='甲', score=80, dur=60000)
        self._add_score(quiz='test_a', openid='oB', nickname='乙', score=90, dur=60000)
        self._add_score(quiz='test_a', openid='oC', nickname='丙', score=90, dur=30000)
        self._add_score(quiz='test_b', openid='oD', nickname='丁', score=100)
        r = self.client.get('/api/mp-admin/quiz-rank?quiz_id=test_a').get_json()
        self.assertTrue(r['success'])
        d = r['data']
        self.assertEqual(d['quiz_id'], 'test_a')
        self.assertEqual(d['name'], 'test_a')  # 无映射时回退 quiz_id
        self.assertEqual(d['total'], 3)
        self.assertEqual([i['nickname'] for i in d['items']], ['丙', '乙', '甲'])
        self.assertEqual(d['items'][0]['score'], 90)
        self.assertEqual(d['items'][0]['duration_ms'], 30000)
        # 每行带分数/满分/测试时间
        for it in d['items']:
            self.assertEqual(it['full_score'], 100)
            self.assertTrue(it['created_at'])
        # 分页:page_size=2 → 第1页2条、共2页
        r2 = self.client.get('/api/mp-admin/quiz-rank?quiz_id=test_a&page_size=2').get_json()
        self.assertEqual(len(r2['data']['items']), 2)
        self.assertEqual(r2['data']['total_pages'], 2)
        r3 = self.client.get('/api/mp-admin/quiz-rank?quiz_id=test_a&page_size=2&page=2').get_json()
        self.assertEqual([i['nickname'] for i in r3['data']['items']], ['甲'])

    def test_rank_auth_and_params(self):
        r = self.client.get('/api/mp-admin/quiz-rank?quiz_id=test_a')
        self.assertEqual(r.status_code, 401)
        self._login()
        r2 = self.client.get('/api/mp-admin/quiz-rank')
        self.assertEqual(r2.status_code, 400)
        # 查无此 quiz:空榜不是错误
        r3 = self.client.get('/api/mp-admin/quiz-rank?quiz_id=nope').get_json()
        self.assertTrue(r3['success'])
        self.assertEqual(r3['data']['items'], [])
        self.assertEqual(r3['data']['total'], 0)


class QuizNameTests(MpAdminTestCase):
    """中文名解析:scores 的 quiz_id 不带 test_ 前缀/驼峰时,也能映射到计数 key 的中文名。"""

    def test_overview_quiz_top_names(self):
        self._login()
        # 客户端实际上报形态:不带 test_ 前缀、可能驼峰
        self._add_score(quiz='pastlifeWho', openid='oA')
        self._add_score(quiz='friend', openid='oB')
        self._add_score(quiz='childIntelligence', openid='oC')
        self._add_score(quiz='test_mbti', openid='oD')   # 带前缀直接命中
        r = self.client.get('/api/mp-admin/overview').get_json()
        names = {q['quiz_id']: q['name'] for q in r['data']['quiz_top']}
        self.assertEqual(names['pastlifeWho'], '前世测试')
        self.assertEqual(names['friend'], '好友印象测试')
        self.assertEqual(names['childIntelligence'], '智力测试')
        self.assertEqual(names['test_mbti'], 'MBTI人格测试')

    def test_scores_and_rank_names(self):
        self._login()
        self._add_score(quiz='pastlifeWho', openid='oA', nickname='甲')
        # 成绩列表带 name 字段
        r = self.client.get('/api/mp-admin/scores?quiz_id=pastlifeWho').get_json()
        self.assertEqual(r['data']['items'][0]['name'], '前世测试')
        # 排行榜接口同样解析
        r2 = self.client.get('/api/mp-admin/quiz-rank?quiz_id=pastlifeWho').get_json()
        self.assertEqual(r2['data']['name'], '前世测试')

    def test_rooms_names(self):
        self._login()
        with m._db() as conn:
            conn.execute(
                "INSERT INTO rooms(room_code, quiz_id, openid_a, state, expires_at, created_at)"
                " VALUES ('1234', 'tool_pkRoom', 'oA', 'waiting', '2099-01-01 00:00:00', ?)",
                (_utcnow_str(),))
        r = self.client.get('/api/mp-admin/rooms').get_json()
        item = next(x for x in r['data']['items'] if x['room_code'] == '1234')
        self.assertEqual(item['name'], '双人默契大作战')


class InjectionGuardTests(MpAdminTestCase):
    """注入防护:复用 Jarvis 攻击模式库,小程序上报字段一律先检测再入库/回显。"""

    def test_score_quiz_id_whitelist(self):
        # quiz_id 白名单:英文标识符形态,含标签/引号/空格一律拒绝
        for evil in ['<script>alert(1)</script>', "test' --", 'test a', 'test\x00']:
            r = self.client.post('/api/mp/game/score', json={
                'openid': 'oA', 'quiz_id': evil, 'score': 80, 'full_score': 100,
                'duration_ms': 60000, 'answer_count': 10})
            self.assertEqual(r.status_code, 200)
            self.assertFalse(r.get_json()['success'], evil)
        # 正常标识符不受影响(驼峰/数字/连字符)
        for ok in ['pastlifeWho', 'test_007', 'tool-x']:
            self.assertTrue(m.QUIZ_ID_RE.match(ok), ok)

    def test_score_nickname_attack_rejected(self):
        # 昵称含攻击特征:入库前拦截,回退默认昵称(不依赖微信检测,本地先拦)
        r = self.client.post('/api/mp/game/score', json={
            'openid': 'oA', 'quiz_id': 'test_a', 'nickname': '<img src=x onerror=alert(1)>',
            'score': 80, 'full_score': 100, 'duration_ms': 60000, 'answer_count': 10})
        data = r.get_json()
        self.assertTrue(data['success'])
        self.assertEqual(data['nickname'], m.DEFAULT_NICKNAME)
        with m._db() as conn:
            row = conn.execute('SELECT nickname FROM scores').fetchone()
            self.assertEqual(row['nickname'], m.DEFAULT_NICKNAME)

    def test_stat_name_attack_dropped(self):
        # stat_names 的 name 能覆盖内置映射并回显后台:攻击特征直接丢弃,计数照常+1
        self.client.post('/api/mp/game/stat/inc', json={
            'key': 'test_a', 'openid': 'oA', 'name': '<svg onload=alert(1)>'})
        with m._db() as conn:
            cnt = conn.execute("SELECT count FROM stats WHERE key='test_a'").fetchone()
            nm = conn.execute("SELECT name FROM stat_names WHERE key='test_a'").fetchone()
            self.assertEqual(cnt['count'], 1)          # 计数不受影响
            self.assertIsNone(nm)                       # 攻击名未入库
        # 正常中文名照常入库
        self.client.post('/api/mp/game/stat/inc', json={
            'key': 'test_b', 'openid': 'oA', 'name': '正常测试名'})
        with m._db() as conn:
            nm = conn.execute("SELECT name FROM stat_names WHERE key='test_b'").fetchone()
            self.assertEqual(nm['name'], '正常测试名')

    def test_legacy_dirty_nickname_sanitized_on_echo(self):
        # 存量脏数据(防护上线前入库):玩家排行榜与后台列表回显时兜底替换
        self._add_score(quiz='test_a', openid='oA',
                        nickname='<img src=x onerror=alert(1)>')
        # 玩家侧排行榜 TOP10
        r = self.client.get('/api/mp/game/rank?quiz_id=test_a&openid=oA').get_json()
        self.assertEqual(r['top'][0]['nickname'], m.DEFAULT_NICKNAME)
        # 后台成绩列表
        self._login()
        r2 = self.client.get('/api/mp-admin/scores?quiz_id=test_a').get_json()
        self.assertEqual(r2['data']['items'][0]['nickname'], m.DEFAULT_NICKNAME)
        # 后台排行榜弹窗
        r3 = self.client.get('/api/mp-admin/quiz-rank?quiz_id=test_a').get_json()
        self.assertEqual(r3['data']['items'][0]['nickname'], m.DEFAULT_NICKNAME)

    def test_room_quiz_id_whitelist(self):
        r = self.client.post('/api/mp/game/room', json={
            'openid': 'oA', 'quiz_id': '<script>'})
        self.assertFalse(r.get_json()['success'])

    def test_openid_rejected_at_submission(self):
        # openid 也会入库并回显后台(玩家列):所有提交入口统一白名单拦截
        evil = '<script>alert(1)</script>'
        # 成绩上报
        r1 = self.client.post('/api/mp/game/score', json={
            'openid': evil, 'quiz_id': 'test_a', 'score': 80, 'full_score': 100,
            'duration_ms': 60000, 'answer_count': 10})
        self.assertFalse(r1.get_json()['success'])
        # 创建/加入房间、交卷、轮询、逐题对照
        self.assertFalse(self.client.post('/api/mp/game/room', json={
            'openid': evil, 'quiz_id': 'pk'}).get_json()['success'])
        self.assertFalse(self.client.post('/api/mp/game/room/join', json={
            'room_code': '1234', 'openid': evil}).get_json()['success'])
        self.assertFalse(self.client.post('/api/mp/game/room/answer', json={
            'room_code': '1234', 'openid': evil, 'answers': 'AB'}).get_json()['success'])
        self.assertFalse(self.client.get(
            '/api/mp/game/room/status?room_code=1234&openid=' + evil).get_json()['success'])
        self.assertFalse(self.client.get(
            '/api/mp/game/room/detail?room_code=1234&openid=' + evil).get_json()['success'])
        # 计数上报(error 字段规范)
        self.assertFalse(self.client.post('/api/mp/game/stat/inc', json={
            'key': 'test_a', 'openid': evil}).get_json()['success'])
        # 排行榜查询
        self.assertFalse(self.client.get(
            '/api/mp/game/rank?quiz_id=test_a&openid=' + evil).get_json()['success'])
        # 库里绝无脏 openid
        with m._db() as conn:
            self.assertIsNone(conn.execute(
                'SELECT * FROM scores WHERE openid = ?', (evil,)).fetchone())

    def test_room_code_rejected_at_submission(self):
        # 房间码只认4位数字:标签/超长/字母一律拒
        for evil in ['<img>', '12345', '12a4', "1' --"]:
            r = self.client.post('/api/mp/game/room/join', json={
                'room_code': evil, 'openid': 'oA'})
            self.assertFalse(r.get_json()['success'], evil)


class StatsTests(MpAdminTestCase):
    def test_list_save_delete(self):
        self._login()
        with m._db() as conn:
            conn.execute("INSERT INTO stats(key, count) VALUES ('test_a', 3)")
        r = self.client.get('/api/mp-admin/stats').get_json()
        self.assertEqual(r['data']['total'], 1)
        self.assertEqual(r['data']['items'][0]['key'], 'test_a')
        self.assertEqual(r['data']['sum'], 3)

    def test_list_filter_by_type(self):
        # ?type=test|tool 按前缀过滤:分开查看测试/工具的使用情况
        self._login()
        with m._db() as conn:
            conn.execute("INSERT INTO stats(key, count) VALUES ('test_a', 3)")
            conn.execute("INSERT INTO stats(key, count) VALUES ('test_b', 4)")
            conn.execute("INSERT INTO stats(key, count) VALUES ('tool_x', 5)")
            conn.execute("INSERT INTO stats(key, count) VALUES ('other', 9)")
            conn.execute("INSERT INTO stats(key, count) VALUES ('article_x', 6)")
        r_test = self.client.get('/api/mp-admin/stats?type=test').get_json()
        keys = [i['key'] for i in r_test['data']['items']]
        self.assertEqual(keys, ['test_b', 'test_a'])       # 按次数降序
        self.assertEqual(r_test['data']['sum'], 7)
        r_tool = self.client.get('/api/mp-admin/stats?type=tool').get_json()
        self.assertEqual([i['key'] for i in r_tool['data']['items']], ['tool_x'])
        self.assertEqual(r_tool['data']['sum'], 5)
        r_art = self.client.get('/api/mp-admin/stats?type=article').get_json()
        self.assertEqual([i['key'] for i in r_art['data']['items']], ['article_x'])
        self.assertEqual(r_art['data']['sum'], 6)
        # 非法 type 等于不过滤
        r_all = self.client.get('/api/mp-admin/stats?type=xyz').get_json()
        self.assertEqual(r_all['data']['total'], 5)
        # 新增(驼峰自动归一) + 缓存失效
        self.client.get('/api/mp/game/stat?keys=test_new')  # 0 进缓存
        r2 = self.client.post('/api/mp-admin/stats/save',
                              json={'password': 'pw', 'key': 'Test_New', 'count': 7}).get_json()
        self.assertTrue(r2['success'])
        q = self.client.get('/api/mp/game/stat?keys=test_new').get_json()
        self.assertEqual(q['counts']['test_new'], 7)
        # 非法
        bad = self.client.post('/api/mp-admin/stats/save',
                               json={'password': 'pw', 'key': '坏key', 'count': 1})
        self.assertEqual(bad.status_code, 400)
        # 删除
        r3 = self.client.post('/api/mp-admin/stats/delete',
                              json={'password': 'pw', 'key': 'test_new'}).get_json()
        self.assertTrue(r3['success'])
        q2 = self.client.get('/api/mp/game/stat?keys=test_new').get_json()
        self.assertEqual(q2['counts']['test_new'], 0)
        r4 = self.client.post('/api/mp-admin/stats/delete',
                              json={'password': 'pw', 'key': 'test_new'})
        self.assertEqual(r4.status_code, 404)

    def test_echohug_keys_hidden_from_admin(self):
        # 弹幕墙「抱抱」走 /stat/inc 的 echohug_<留言id> 技术计数:
        # 每条留言一个key,量大且非业务内容 → 不进热力图/计数列表/导出条目口径
        self._login()
        self.client.post('/api/mp/game/stat/inc',
                         json={'key': 'echohug_7', 'openid': 'oA'})
        d = self.client.get('/api/mp-admin/overview').get_json()['data']
        self.assertNotIn('echohug_7', {i['key'] for i in d['hourly_top']['items']})
        self.assertNotIn('echohug_7', {i['key'] for i in d['daily_top']['items']})
        items = self.client.get('/api/mp-admin/stats').get_json()['data']['items']
        self.assertNotIn('echohug_7', {i['key'] for i in items})


class RoomsTests(MpAdminTestCase):
    def test_list_and_delete(self):
        self._login()
        with m._db() as conn:
            conn.execute('''INSERT INTO rooms
                (room_code, quiz_id, openid_a, answers_a, openid_b, answers_b, state,
                 answer_count, match_percent, expires_at, created_at)
                VALUES ('1234', 'pk', 'oA', 'AB', 'oB', 'AC', 'finished', 2, 50,
                        '2099-01-01 00:00:00', ?)''', (_utcnow_str(),))
        r = self.client.get('/api/mp-admin/rooms').get_json()
        self.assertEqual(r['data']['total'], 1)
        item = r['data']['items'][0]
        self.assertEqual(item['players'], 2)
        self.assertTrue(item['a_submitted'] and item['b_submitted'])
        self.assertEqual(item['match_percent'], 50)
        # 按状态过滤
        r2 = self.client.get('/api/mp-admin/rooms?state=waiting').get_json()
        self.assertEqual(r2['data']['total'], 0)
        # 删除
        r3 = self.client.post('/api/mp-admin/rooms/delete',
                              json={'password': 'pw', 'room_code': '1234'}).get_json()
        self.assertTrue(r3['success'])
        r4 = self.client.post('/api/mp-admin/rooms/delete',
                              json={'password': 'pw', 'room_code': '1234'})
        self.assertEqual(r4.status_code, 404)

    def test_expired_display_and_days_window(self):
        self._login()
        ten_days_ago = (datetime.now(timezone.utc) - timedelta(days=10)) \
            .strftime('%Y-%m-%d %H:%M:%S')
        with m._db() as conn:
            # 10天前建的活跃房间:默认近7天不显示,days=0 显示
            conn.execute('''INSERT INTO rooms(room_code, quiz_id, openid_a, state, expires_at, created_at)
                VALUES ('1111', 'pk', 'oA', 'waiting', '2099-01-01 00:00:00', ?)''',
                (ten_days_ago,))
            # 已过期的房间:统一展示为 expired
            conn.execute('''INSERT INTO rooms(room_code, quiz_id, openid_a, state, expires_at, created_at)
                VALUES ('2222', 'pk', 'oB', 'waiting', '2000-01-01 00:00:00', ?)''',
                (_utcnow_str(),))
        default = self.client.get('/api/mp-admin/rooms').get_json()
        self.assertEqual(default['data']['total'], 1)  # 只有今天建的那个
        self.assertEqual(default['data']['items'][0]['room_code'], '2222')
        self.assertEqual(default['data']['items'][0]['state'], 'expired')
        all_days = self.client.get('/api/mp-admin/rooms?days=0').get_json()
        self.assertEqual(all_days['data']['total'], 2)


class EchoAdminTests(MpAdminTestCase):
    """弹幕墙留言管理:补录/列表(实时抱抱数)/删除(连带清 echohug_ 计数)。"""

    def test_list_create_delete(self):
        self._login()
        r = self.client.post('/api/mp-admin/echo/create', json={
            'password': 'pw', 'wall_id': '20260906', 'text': '管理员补录', 'nickname': '小编'
        }).get_json()
        self.assertTrue(r['success'])
        admin_id = r['id']
        # 补录辱骂被本地词库拦(管理员内容不走微信送检,但本地硬底线照拦)
        bad = self.client.post('/api/mp-admin/echo/create', json={
            'password': 'pw', 'wall_id': '20260906', 'text': '你就是个傻逼'})
        self.assertEqual(bad.status_code, 400)
        # 玩家留言直接入库(绕开需要微信送检的 /game/echo);三人各抱一次=3
        with m._db() as conn:
            cur = conn.execute(
                "INSERT INTO echo_wall(wall_id, openid, nickname, text, hugs, ts) "
                "VALUES ('20260906', 'oA', '甲', '玩家留言', 0, ?)", (int(time.time()),))
            player_id = cur.lastrowid
        for oid in ('o1', 'o2', 'o3'):
            self.client.post('/api/mp/game/stat/inc',
                             json={'key': f'echohug_{player_id}', 'openid': oid})
        # 同人重复抱:永久去重(清掉10秒防刷窗口,证明不是窗口吞的),计数不涨
        m._stat_dedup.clear()
        rep = self.client.post('/api/mp/game/stat/inc',
                               json={'key': f'echohug_{player_id}', 'openid': 'o1'}).get_json()
        self.assertEqual(rep['count'], 3)
        d = self.client.get('/api/mp-admin/echo?wall_id=20260906').get_json()['data']
        self.assertEqual(d['total'], 2)
        by_id = {i['id']: i for i in d['items']}
        self.assertEqual(by_id[player_id]['hugs'], 3)     # 实时抱抱数回填
        self.assertEqual(by_id[player_id]['openid'], 'oA')
        self.assertEqual(by_id[admin_id]['nickname'], '小编')
        self.assertEqual(self.client.get('/api/mp-admin/echo?days=0')
                         .get_json()['data']['total'], 2)
        # 删除:连带清 echohug_ 计数,不留孤儿条目
        r2 = self.client.post('/api/mp-admin/echo/delete',
                              json={'password': 'pw', 'id': player_id}).get_json()
        self.assertTrue(r2['success'])
        with m._db() as conn:
            self.assertIsNone(conn.execute('SELECT * FROM echo_wall WHERE id = ?',
                                           (player_id,)).fetchone())
            self.assertIsNone(conn.execute('SELECT * FROM stats WHERE key = ?',
                                           (f'echohug_{player_id}',)).fetchone())
            self.assertIsNone(conn.execute('SELECT * FROM echo_hugs WHERE hug_key = ?',
                                           (f'echohug_{player_id}',)).fetchone())
        r3 = self.client.post('/api/mp-admin/echo/delete',
                              json={'password': 'pw', 'id': player_id})
        self.assertEqual(r3.status_code, 404)
        # 删除后立刻再查:剩余留言照常返回(前端删除后自动刷新依赖这一点,不能空)
        d2 = self.client.get('/api/mp-admin/echo?wall_id=20260906').get_json()['data']
        self.assertEqual([i['id'] for i in d2['items']], [admin_id])
        self.assertEqual(d2['total'], 1)
        # 分页越界:页码超出总页数时返回空items但不报错(前端负责回退到最后一页)
        d3 = self.client.get('/api/mp-admin/echo?page=9').get_json()['data']
        self.assertEqual(d3['items'], [])
        self.assertEqual(d3['total'], 1)

    def test_echo_write_requires_admin_login(self):
        r0 = self.client.post('/api/mp-admin/echo/create',
                              json={'wall_id': 'w', 'text': 'x'})
        self.assertEqual(r0.status_code, 401)             # 未登录
        self._login()
        r1 = self.client.post('/api/mp-admin/echo/create',
                              json={'wall_id': 'w2', 'text': 'x'})   # 无密码字段
        self.assertEqual(r1.status_code, 200)
        r2 = self.client.get('/api/mp-admin/echo')
        self.assertEqual(r2.status_code, 200)             # 只读有登录即可

    def test_seed_import_once_and_badge(self):
        # 冷启动种子:一键导入20条(is_seed=1+预设抱抱数+openid='seed'),
        # 全局幂等(换墙重导也拒绝);后台列表带 is_seed 与预设 hugs
        r0 = self.client.post('/api/mp-admin/echo/seed', json={'wall_id': 'w'})
        self.assertEqual(r0.status_code, 401)             # 未登录不可导入
        self._login()
        r = self.client.post('/api/mp-admin/echo/seed',
                             json={'wall_id': '20260906'}).get_json()
        self.assertTrue(r['success'])
        self.assertEqual(r['count'], 20)
        d = self.client.get('/api/mp-admin/echo?wall_id=20260906').get_json()['data']
        self.assertEqual(d['total'], 20)
        self.assertTrue(all(i['is_seed'] for i in d['items']))
        self.assertEqual({i['openid'] for i in d['items']}, {'seed'})
        self.assertIn(88, {i['hugs'] for i in d['items']})   # 预设抱抱数带出
        self.assertTrue(all(i['category'] in m.ECHO_CATEGORIES for i in d['items']))
        self.assertIn('职场', {i['category'] for i in d['items']})  # 种子分类带出(加班那条)
        # 玩家侧列表同墙可见(昵称/抱抱数照常展示,不带 is_seed 字段)
        gl = self.client.get('/api/mp/game/echo?wall_id=20260906').get_json()
        self.assertEqual(len(gl['list']), 20)
        self.assertTrue(all('is_seed' not in x for x in gl['list']))
        # 幂等:重复导入(哪怕换墙)拒绝,数据不翻倍
        r2 = self.client.post('/api/mp-admin/echo/seed', json={'wall_id': '20260907'})
        self.assertEqual(r2.status_code, 409)
        self.assertEqual(self.client.get('/api/mp-admin/echo?days=0')
                         .get_json()['data']['total'], 20)
        # 非法 wall_id
        bad = self.client.post('/api/mp-admin/echo/seed', json={'wall_id': 'bad wall'})
        self.assertEqual(bad.status_code, 400)


class EchoBanTests(MpAdminTestCase):
    """弹幕墙用户封禁:封禁/解封/列表 + 发布接口拦截(永久/临时/过期懒失效)。"""

    def _allow_sec(self):
        # 打桩微信送检,让玩家侧 /game/echo 正常走通
        m._do_msg_sec_check = lambda openid, text, scene=1: {
            'errcode': 0, 'result': {'suggest': 'pass', 'label': 100}}

    def _post_echo(self, openid):
        return self.client.post('/api/mp/game/echo',
                                json={'openid': openid, 'wall_id': '20260906',
                                      'text': '留言内容'}).get_json()

    def test_ban_unban_and_enforcement(self):
        self._allow_sec()
        r0 = self.client.post('/api/mp-admin/echo/ban', json={'openid': 'oBad'})
        self.assertEqual(r0.status_code, 401)             # 未登录不可封禁
        self._login()
        self.assertTrue(self._post_echo('oGood')['success'])   # 未封禁者照常发
        # 永久封禁 → 发布被拒;抱抱计数(stat/inc)不受影响
        self.assertTrue(self.client.post('/api/mp-admin/echo/ban',
                                         json={'openid': 'oBad', 'reason': '刷屏'}).get_json()['success'])
        rep = self._post_echo('oBad')
        self.assertFalse(rep['success'])
        self.assertIn('封禁', rep['message'])
        self.assertTrue(self.client.post('/api/mp/game/stat/inc',
                                         json={'key': 'echohug_1', 'openid': 'oBad'}).get_json()['success'])
        # 封禁列表:permanent 状态 + 原因带出
        bans = self.client.get('/api/mp-admin/echo/bans').get_json()['data']
        self.assertEqual(bans['total'], 1)
        self.assertEqual(bans['items'][0]['status'], 'permanent')
        self.assertEqual(bans['items'][0]['reason'], '刷屏')
        # 临时封禁覆盖永久(状态 active),解封后恢复发布;重复解封幂等
        self.assertTrue(self.client.post('/api/mp-admin/echo/ban',
                                         json={'openid': 'oBad', 'days': 7}).get_json()['success'])
        bans = self.client.get('/api/mp-admin/echo/bans').get_json()['data']
        self.assertEqual(bans['items'][0]['status'], 'active')
        self.assertTrue(self.client.post('/api/mp-admin/echo/unban',
                                         json={'openid': 'oBad'}).get_json()['success'])
        self.assertTrue(self._post_echo('oBad')['success'])
        self.assertTrue(self.client.post('/api/mp-admin/echo/unban',
                                         json={'openid': 'oBad'}).get_json()['success'])

    def test_expired_ban_allows_post(self):
        self._allow_sec()
        self._login()
        now = int(time.time())
        with m._db() as conn:   # 直接种一条已过期的临时封禁
            conn.execute('INSERT INTO echo_bans(openid, reason, expires_at, created_at) '
                         "VALUES ('oOld', '历史', ?, ?)", (now - 10, now - 86400))
        self.assertTrue(self._post_echo('oOld')['success'])   # 过期懒失效,可正常发布
        bans = self.client.get('/api/mp-admin/echo/bans').get_json()['data']
        self.assertEqual(bans['items'][0]['status'], 'expired')

    def test_ban_param_validation(self):
        self._login()
        self.assertEqual(self.client.post('/api/mp-admin/echo/ban',
                                          json={'openid': 'bad openid!'}).status_code, 400)
        self.assertEqual(self.client.post('/api/mp-admin/echo/ban',
                                          json={'openid': 'oX', 'days': 4000}).status_code, 400)
        self.assertEqual(self.client.post('/api/mp-admin/echo/ban',
                                          json={'openid': 'oX', 'days': 'abc'}).status_code, 400)


class CleanupTests(MpAdminTestCase):
    def test_cleanup_dev_data(self):
        self._login()
        with m._db() as conn:
            # 遗留 quiz + 联调 openid 的成绩
            for quiz, oid in (('selftest_tmp', 'oReal'), ('wealth', 'oReal'),
                              ('test_a', 'vkself001'), ('test_a', 'otest001'),
                              ('test_a', 'oReal')):
                conn.execute('''INSERT INTO scores(quiz_id, openid, nickname, score,
                    full_score, duration_ms, created_at) VALUES (?, ?, ?, 1, 100, 60000, ?)''',
                    (quiz, oid, 'n', _utcnow_str()))
            # 联调 openid 的房间
            conn.execute('''INSERT INTO rooms(room_code, quiz_id, openid_a, state, expires_at, created_at)
                VALUES ('9001', 'pk', 'vkself001', 'waiting', '2099-01-01 00:00:00', ?)''',
                (_utcnow_str(),))
            # 计数:test_a=10,探针=4;流水:联调贡献3次、真人贡献2次
            conn.execute("INSERT INTO stats(key, count) VALUES ('test_a', 10)")
            conn.execute("INSERT INTO stats(key, count) VALUES ('tool_probe_diag', 4)")
            for oid in ('vkself001', 'vkself002', 'otest001', 'oReal', 'oReal2'):
                conn.execute("INSERT INTO stat_log(key, ts, openid) VALUES ('test_a', ?, ?)",
                             (int(datetime.now(timezone.utc).timestamp()), oid))
            conn.execute("INSERT INTO stat_log(key, ts, openid) VALUES ('tool_probe_diag', 123, 'oReal')")
        r = self.client.post('/api/mp-admin/cleanup-dev-data',
                             json={'password': 'pw'}).get_json()
        self.assertTrue(r['success'])
        res = r['result']
        self.assertEqual(res['scores'], 4)            # selftest_tmp+wealth+vkself+otest
        self.assertEqual(res['rooms'], 1)
        self.assertEqual(res['log_rows'], 3)          # vkself×2 + otest×1
        self.assertEqual(res['decremented'], {'test_a': 3})
        self.assertTrue(res['probe_cleared'])
        with m._db() as conn:
            left_scores = conn.execute('SELECT COUNT(*) AS c FROM scores').fetchone()['c']
            left_logs = conn.execute('SELECT COUNT(*) AS c FROM stat_log').fetchone()['c']
            test_a = conn.execute("SELECT count FROM stats WHERE key='test_a'").fetchone()['count']
            probe = conn.execute("SELECT count FROM stats WHERE key='tool_probe_diag'").fetchone()['count']
        self.assertEqual(left_scores, 1)              # 只剩 oReal 的 test_a
        self.assertEqual(left_logs, 2)                # 真人的2条 + 探针日志已删
        self.assertEqual(test_a, 7)                   # 10 - 联调3次
        self.assertEqual(probe, 0)
        # 可重复执行:再跑一遍不报错、不再删东西
        r2 = self.client.post('/api/mp-admin/cleanup-dev-data',
                              json={'password': 'pw'}).get_json()
        self.assertTrue(r2['success'])
        self.assertEqual(r2['result']['scores'], 0)

    def test_cleanup_requires_login(self):
        r = self.client.post('/api/mp-admin/cleanup-dev-data', json={})
        self.assertEqual(r.status_code, 401)


if __name__ == '__main__':
    unittest.main()
